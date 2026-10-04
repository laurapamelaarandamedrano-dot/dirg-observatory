"""Build the reader-friendly Excel companion (one workbook, one guide, everything organized).

Usage: python scripts/make_workbook.py  ->  dist/DIRG_Observatory_v<version>.xlsx
Every value comes from data/public/ and config/; nothing is typed by hand.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.build_site import VERSION  # noqa: E402
from src.common import STATES as STATE_CAT  # noqa: E402

PUB = ROOT / "data" / "public"
CFG = yaml.safe_load((ROOT / "config" / "scenarios.yaml").read_text(encoding="utf-8"))
VARS = yaml.safe_load((ROOT / "config" / "variables.yaml").read_text(encoding="utf-8"))
QA = json.loads((ROOT / "reports" / "quality_report.json").read_text(encoding="utf-8"))

# ------------------------------------------------------------------ palette
NAVY, INK, MUTED, LINE = "10233F", "1F2733", "5B6676", "D5DBE3"
PILLAR = {  # header colour, light tint
    "W": ("2F8FC4", "E3F2FA"), "E": ("C9861C", "FCF0DC"), "D": ("6E5FD0", "ECE9FB"),
    "V": ("2E9A6A", "E2F4EA"), "G": ("8A7E5E", "F3EFE3"), "C": ("6B7685", "EEF0F3"), "I": (NAVY, "DCEFFB"),
}
PILLAR_MID = {"W": "8CC6E6", "E": "F0C77E", "D": "B9AFF0", "V": "8FD1B1"}
TIER = {"high": "BFE3F5", "moderate": "DCE6F0", "limited": "FBE3B0", "insufficient": "E2E2E2"}
LEDGER = {"observed": "D6ECFA", "administrative": "D6ECFA", "documentary": "D6ECFA", "geospatial": "D6ECFA",
          "derived": "E6EBF1", "estimated": "FCEBC4", "modelled": "E8E2FB", "hypothesized": "EFEFEF", "unknown": "EFEFEF"}
F = "Arial"
thin = Side(style="thin", color=LINE)
BORDER = Border(bottom=thin)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def font(size=10, bold=False, color=INK, italic=False, underline=None):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic, underline=underline)


STATES = pd.read_csv(PUB / "territories.csv", dtype={"cve_ent": str})
NAME = {k: v[1] for k, v in STATE_CAT.items()}
obs = pd.read_csv(PUB / "observations.csv", dtype={"territory_id": str})
res = pd.read_csv(PUB / "index_results.csv", dtype={"territory_id": str})
cov = pd.read_csv(PUB / "coverage.csv", dtype={"territory_id": str})
YEARS = CFG["years"]
YEAR_NOTE = {2026: "partial", 2027: "projection"}

wb = Workbook()
ENDS: dict[str, tuple[int, int]] = {}


def sheet_title(ws, title, subtitle, width_cols, band="I"):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = font(15, True, "FFFFFF")
    ws["A2"] = subtitle
    ws["A2"].font = font(10, False, "DCE6F0")
    for r in (1, 2):
        for c in range(1, width_cols + 1):
            ws.cell(r, c).fill = fill(PILLAR[band][0] if band != "I" else NAVY)
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 18


def header_row(ws, row, headers, colours=None, height=32):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row, i, h)
        bg = (colours[i - 1] if colours else NAVY)
        c.fill = fill(bg)
        c.font = font(9, True, "FFFFFF")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
    ws.row_dimensions[row].height = height


def write_table(ws, df, start_row, colours=None, number_formats=None, wrap_cols=(), widths=None, zebra=True):
    header_row(ws, start_row, list(df.columns), colours)
    for r_i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for c_i, v in enumerate(row, 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            cell = ws.cell(r_i, c_i, v)
            cell.font = font(9)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(c_i in wrap_cols))
            if zebra and (r_i - start_row) % 2 == 0:
                cell.fill = fill("F7F9FB")
            if number_formats and c_i in number_formats:
                cell.number_format = number_formats[c_i]
    end = start_row + len(df)
    ENDS[ws.title] = (start_row + 1, end)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(df.columns))}{end}"
    ws.freeze_panes = ws.cell(start_row + 1, 1)
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return end


def note(ws, row, text, cols=10, colour=MUTED, italic=True):
    c = ws.cell(row, 1, text)
    c.font = font(9, False, colour, italic)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    total = sum((ws.column_dimensions[get_column_letter(i)].width or 9) for i in range(1, cols + 1))
    lines = -(-len(text) * 1.15 // max(total, 20))
    ws.row_dimensions[row].height = max(15, 14 * lines + 4)


# ================================================================== Guide
g = wb.active
g.title = "Guide"
g.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGH", (36, 70, 10, 10, 10, 10, 10, 10)):
    g.column_dimensions[col].width = w
for r in range(1, 5):
    for c in range(1, 9):
        g.cell(r, c).fill = fill(NAVY)
g["A2"] = "Digital Infrastructure & Resource Governance Observatory"
g["A2"].font = font(18, True, "FFFFFF")
g["A3"] = f"Mexico · 32 states · 2020 to 2027 · Release {VERSION} · data as of {CFG['as_of_date']}"
g["A3"].font = font(10, False, "BFD3E8")
g.row_dimensions[2].height = 30

r = 6
g.cell(r, 1, "What this workbook is").font = font(12, True, NAVY)
r += 1
note(g, r, "The complete public dataset of the observatory in one place. It shows where digital infrastructure, water pressure, energy systems and resource-governance instruments coincide across Mexico's states, year by year. Every value carries its source and how it was obtained, and every state-year carries a data-coverage score.", 8, INK, False)
r += 1
note(g, r, "It describes and models relationships. It does not establish causation, legal violations or the responsibility of any company. No company names or facility locations are included.", 8, "8A4B0F", False)

r += 2
g.cell(r, 1, "Start here").font = font(12, True, NAVY)
steps = [
    ("1  SRTI by year", "A one-glance matrix: each state's tension index for every year. Darker = higher."),
    ("2  Pillars", "Why a state scores what it scores: the four pillars, its rank, the rank range under other weightings, and data coverage."),
    ("3  Indicators", "The actual measured values behind the pillars (drought, aquifers, capacity, cloud regions, population...)."),
    ("4  Variables + Sources", "What every column means, and where every number comes from."),
]
for a, b in steps:
    r += 1
    c = g.cell(r, 1, a); c.font = font(10, True, INK)
    c2 = g.cell(r, 2, b); c2.font = font(10, False, INK); c2.alignment = Alignment(wrap_text=True)
    g.row_dimensions[r].height = 28

r += 2
g.cell(r, 1, "Sheets in this workbook").font = font(12, True, NAVY)
r += 1
for i, h in enumerate(["Sheet", "What you will find", "Rows"], 1):
    c = g.cell(r, i, h); c.fill = fill(NAVY); c.font = font(9, True, "FFFFFF")
sheet_map = [
    ("SRTI by year", "State × year matrix of the Sociotechnical Resource Tension Index (SRTI) and the Resource Baseline Index (RBI, without the digital pillar)", "I"),
    ("Pillars", "Pillar scores (0 to 100), SRTI, RBI, rank, Monte Carlo ranges and coverage for every state-year", "I"),
    ("Indicators", "Raw indicator values, state × year, grouped by pillar colour", "I"),
    ("Scenarios", "The composite under 9 alternative specifications", "I"),
    ("Sensitivity", "How closely each specification agrees with the baseline ranking (Spearman correlation)", "I"),
    ("Coverage", "Data-coverage score and tier per state-year, and how many indicators are known, derived, estimated, modelled or unknown", "I"),
    ("Variables", "Data dictionary: definition, unit, direction, status and source of each variable", "I"),
    ("Sources", "Source registry with licences and what may be redistributed", "I"),
    ("Cloud register", "Documented hyperscale cloud regions (no operators, no coordinates)", "D"),
    ("Governance", "Federal legal and institutional instruments, described neutrally", "G"),
    ("Quality checks", "Automated validation results and ingestion exclusions", "I"),
    ("Observations", "Full long table: one row per state × year × variable, with status and flags", "I"),
]
first_map_row = r + 1
for name, desc, band in sheet_map:
    r += 1
    c = g.cell(r, 1, name)
    c.hyperlink = f"#'{name}'!A1"
    c.font = font(10, True, PILLAR[band][0], underline="single")
    d = g.cell(r, 2, desc); d.font = font(9, False, INK); d.alignment = Alignment(wrap_text=True, vertical="top")
    g.cell(r, 3).font = font(9, False, MUTED)
    g.row_dimensions[r].height = 26
    for cc in range(1, 4):
        g.cell(r, cc).border = BORDER
last_map_row = r
ROWCOUNT_ROWS = {name: first_map_row + i for i, (name, _, _) in enumerate(sheet_map)}

r += 2
g.cell(r, 1, "Reading the colours").font = font(12, True, NAVY)
r += 1
for key, lab in (("W", "Drought and aquifers without availability"), ("E", "Authorized generation capacity and its combustion share"),
                 ("D", "Hyperscale cloud regions in operation"), ("V", "Population density, marginalization, drought vulnerability"),
                 ("G", "Groundwater ordinances, climate-policy instruments (never in the index)"), ("C", "Population and other context (never in the index)")):
    c = g.cell(r, 1, VARS["pillars"][key]["en"])
    c.fill = fill(PILLAR[key][0]); c.font = font(9, True, "FFFFFF")
    g.cell(r, 2, lab).font = font(9, False, INK)
    r += 1
r += 1
g.cell(r, 1, "How each value was obtained").font = font(10, True, INK)
status_text = [
    ("administrative", "Recorded by an official body (e.g. CONAGUA aquifer availability)"),
    ("documentary", "Taken from a document: provider release notes, press, legal texts"),
    ("geospatial", "Computed from official maps (e.g. area under groundwater vedas)"),
    ("derived", "Calculated by this project from recorded data (e.g. yearly drought average)"),
    ("estimated", "Uses planned dates or partial information (e.g. 2026–2027 capacity)"),
    ("modelled", "Comes from a model or projection (e.g. CONAPO population projections)"),
    ("unknown", "No value. Never treated as zero"),
]
for st, txt in status_text:
    r += 1
    c = g.cell(r, 1, st); c.fill = fill(LEDGER[st]); c.font = font(9, True, INK)
    g.cell(r, 2, txt).font = font(9, False, INK)
r += 2
g.cell(r, 1, "Data coverage tiers").font = font(10, True, INK)
for t, txt in (("high", "Coverage ≥ 0.75"), ("moderate", "0.50 to 0.75"), ("limited", "0.25 to 0.50"), ("insufficient", "Below 0.25: no index value is published")):
    r += 1
    c = g.cell(r, 1, t); c.fill = fill(TIER[t]); c.font = font(9, True, INK)
    g.cell(r, 2, txt).font = font(9, False, INK)

r += 2
g.cell(r, 1, "Quick facts (calculated from the sheets)").font = font(12, True, NAVY)
facts_row = r + 1
r += 1
FACT_LABELS = 7
y_last = 2025
srti_col = get_column_letter(2 + YEARS.index(y_last))
r += FACT_LABELS

r += 1
g.cell(r, 1, "How to cite").font = font(12, True, NAVY)
r += 1
note(g, r, f"Aranda Medrano, L. P. (2026). Digital Infrastructure & Resource Governance Observatory: Mexico, 2020–2027 (Version {VERSION}) [Data set and software]. Harvard Dataverse.", 8, INK, False)
r += 1
note(g, r, "Licence: derived data CC BY 4.0; code MIT. Source data keep their own licences (sheet Sources). Methodology, data dictionary and governance audit are in the accompanying documents and the GitHub repository.", 8, MUTED, True)

# ================================================================== SRTI by year
ws = wb.create_sheet("SRTI by year")
ncols = 2 + len(YEARS)
sheet_title(ws, "SRTI by year", "Sociotechnical Resource Tension Index, baseline (equal weights), 0 to 100. Darker = more coincident pressure.", ncols * 2 + 1)
base = res[res.scenario == "baseline"]
rbi = res[res.scenario == "rbi"]
hdr = ["State"] + [f"{y}{' ' + YEAR_NOTE[y] if y in YEAR_NOTE else ''}" for y in YEARS] + [f"Rank {y_last}"]
note(ws, 3, "2026 combines partial-year observations. 2027 is a projection computed on 3 pillars (digital infrastructure cannot be observed yet), so it is not directly comparable with earlier years.", ncols)
header_row(ws, 5, hdr, [NAVY] * len(hdr))
mat = base.pivot(index="territory_id", columns="year", values="value")
order = mat[y_last].sort_values(ascending=False).index
for i, tid in enumerate(order, 6):
    ws.cell(i, 1, NAME[tid]).font = font(9, True)
    for j, y in enumerate(YEARS, 2):
        v = mat.loc[tid, y]
        c = ws.cell(i, j, None if pd.isna(v) else round(float(v), 1)); c.number_format = "0.0"; c.font = font(9)
        if y == 2027:
            c.font = font(9, italic=True)
    c = ws.cell(i, ncols, f"=RANK({srti_col}{i},{srti_col}$6:{srti_col}$37,0)"); c.font = font(9, True, NAVY)
end = 6 + len(order) - 1
ENDS["SRTI by year"] = (6, end)
ws.conditional_formatting.add(f"B6:{get_column_letter(1 + len(YEARS))}{end}",
                              ColorScaleRule(start_type="num", start_value=0, start_color="F4F8FC", mid_type="num", mid_value=40, mid_color="B4D0E8", end_type="num", end_value=80, end_color="5B93C4"))
# RBI block to the right
off = ncols + 2
ws.cell(4, off, "Resource Baseline Index (RBI): same, without the digital pillar").font = font(10, True, PILLAR["V"][0])
for k, h in enumerate(["State"] + [str(y) for y in YEARS], 0):
    c = ws.cell(5, off + k, h); c.fill = fill(PILLAR["V"][0]); c.font = font(9, True, "FFFFFF")
mat2 = rbi.pivot(index="territory_id", columns="year", values="value")
for i, tid in enumerate(order, 6):
    ws.cell(i, off, NAME[tid]).font = font(9)
    for j, y in enumerate(YEARS, 1):
        v = mat2.loc[tid, y]
        c = ws.cell(i, off + j, None if pd.isna(v) else round(float(v), 1)); c.number_format = "0.0"; c.font = font(9)
ws.conditional_formatting.add(f"{get_column_letter(off + 1)}6:{get_column_letter(off + len(YEARS))}{end}",
                              ColorScaleRule(start_type="num", start_value=0, start_color="F3FAF6", mid_type="num", mid_value=40, mid_color="B5DFC8", end_type="num", end_value=80, end_color="5DB389"))
ws.column_dimensions["A"].width = 22
for j in range(2, ncols + 1):
    ws.column_dimensions[get_column_letter(j)].width = 10.5
ws.column_dimensions[get_column_letter(ncols + 1)].width = 3
ws.column_dimensions[get_column_letter(off)].width = 22
for j in range(off + 1, off + 1 + len(YEARS)):
    ws.column_dimensions[get_column_letter(j)].width = 8
ws.freeze_panes = "B6"
note(ws, end + 2, "Sorted by SRTI in 2025. A high value means several pressures coincide in the same state and year; it is not evidence of harm or of any cause. See sheet Pillars for the breakdown and the rank range under alternative weightings.", ncols)

# ================================================================== Pillars
ws = wb.create_sheet("Pillars")
p = base.merge(rbi[["territory_id", "year", "value"]].rename(columns={"value": "RBI"}), on=["territory_id", "year"]).merge(cov[["territory_id", "year", "coverage", "tier"]], on=["territory_id", "year"])
p["State"] = p.territory_id.map(NAME)
df = pd.DataFrame({
    "State": p.State, "Code": p.territory_id, "Year": p.year,
    "Water (W)": p.W, "Energy (E)": p.E, "Digital (D)": p.D, "Vulnerability (V)": p.V,
    "SRTI": p.value, "RBI (no D)": p.RBI, "Rank (1 = highest)": p["rank"],
    "Rank range low (5%)": p.rank_p05, "Rank range high (95%)": p.rank_p95,
    "SRTI range low": p.mc_low, "SRTI range high": p.mc_high,
    "Coverage (0–1)": p.coverage, "Coverage tier": p.tier, "Pillars used": p.n_pillars, "Flags": p["flags"].fillna(""),
}).sort_values(["Year", "SRTI"], ascending=[True, False])
sheet_title(ws, "Pillars", "Pillar scores on pooled 0–100 scales (2020–2025 reference). Higher = more pressure. Ranges come from 2,000 random weightings.", 18)
note(ws, 3, "Empty pillar = not enough data to compute it (never zero). 'Pillars used' below 4 means weights were renormalized over the available pillars and the row is flagged.", 18)
cols = [NAVY, NAVY, NAVY, PILLAR["W"][0], PILLAR["E"][0], PILLAR["D"][0], PILLAR["V"][0], NAVY, PILLAR["V"][0]] + [NAVY] * 9
nf = {i: "0.0" for i in range(4, 10)} | {i: "0" for i in (10, 11, 12)} | {13: "0.0", 14: "0.0", 15: "0.00"}
end = write_table(ws, df, 5, cols, nf, widths=[20, 6, 7, 10, 10, 10, 12, 9, 10, 10, 11, 11, 10, 10, 11, 12, 8, 24])
for col, key in zip("DEFG", "WEDV"):
    ws.conditional_formatting.add(f"{col}6:{col}{end}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="num", end_value=100, end_color=PILLAR_MID[key]))
ws.conditional_formatting.add(f"H6:H{end}", ColorScaleRule(start_type="num", start_value=0, start_color="F4F8FC", mid_type="num", mid_value=40, mid_color="B4D0E8", end_type="num", end_value=80, end_color="5B93C4"))
for t, hx in TIER.items():
    ws.conditional_formatting.add(f"P6:P{end}", CellIsRule(operator="equal", formula=[f'"{t}"'], fill=fill(hx)))

# ================================================================== Indicators (wide)
ws = wb.create_sheet("Indicators")
order_vars = [k for p_ in ("W", "E", "D", "V", "C", "G") for k, v in VARS["variables"].items() if v["pillar"] == p_]
wide = obs.pivot(index=["territory_id", "year"], columns="variable_id", values="value")[order_vars].reset_index()
wide.insert(0, "State", wide.territory_id.map(NAME))
wide = wide.rename(columns={"territory_id": "Code", "year": "Year"})
heads = ["State", "Code", "Year"] + [f"{VARS['variables'][k]['label']['en']} ({VARS['variables'][k]['unit']})" for k in order_vars]
wide.columns = heads
sheet_title(ws, "Indicators", "Measured values behind the index, as published by the sources or derived from them. Header colour = pillar; grey/beige = not in the index.", len(heads))
note(ws, 3, "Blank = unknown (not zero). How each value was obtained (status), its vintage and flags are in sheet Observations. Definitions in sheet Variables.", len(heads))
cols = [NAVY] * 3 + [PILLAR[VARS["variables"][k]["pillar"]][0] for k in order_vars]
nf = {i: "#,##0.0" for i in range(4, len(heads) + 1)}
end = write_table(ws, wide, 5, cols, nf, widths=[20, 6, 7] + [14] * len(order_vars))
ws.row_dimensions[5].height = 58
for i, k in enumerate(order_vars, 4):
    if k in ("c_population",):
        for rr in range(6, end + 1):
            ws.cell(rr, i).number_format = "#,##0"
    if k in ("d_cloud_regions", "d_cloud_regions_announced", "g_climate_instruments"):
        for rr in range(6, end + 1):
            ws.cell(rr, i).number_format = "0"

# ================================================================== Scenarios
ws = wb.create_sheet("Scenarios")
sc = res.pivot(index=["territory_id", "year"], columns="scenario", values="value").reset_index()
labels = {k: v["label"] for k, v in CFG["scenarios"].items()}
keys = list(CFG["scenarios"])
sc = sc[["territory_id", "year"] + keys]
sc.insert(0, "State", sc.territory_id.map(NAME))
sc.columns = ["State", "Code", "Year"] + [labels[k] for k in keys]
sheet_title(ws, "Scenarios", "The composite under every specification tested. Large differences across a row mean the result depends on methodological choices.", len(sc.columns))
wrow = ["", "", "Weights W / E / D / V →"] + [" / ".join(f"{CFG['scenarios'][k]['weights'][p_]:.2f}" for p_ in "WEDV") + f" · {CFG['scenarios'][k]['aggregation']} · {CFG['scenarios'][k]['normalization'].replace('_pooled', '')}" for k in keys]
for i, v in enumerate(wrow, 1):
    c = ws.cell(4, i, v); c.font = font(8, False, MUTED, True); c.alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[4].height = 40
end = write_table(ws, sc, 5, [NAVY] * 3 + [PILLAR["V"][0] if k == "rbi" else NAVY for k in keys], {i: "0.0" for i in range(4, 4 + len(keys))}, widths=[20, 6, 7] + [14] * len(keys))
ws.row_dimensions[5].height = 44

# ================================================================== Sensitivity
ws = wb.create_sheet("Sensitivity")
sens = pd.read_csv(PUB / "sensitivity.csv")
m = sens.pivot(index="scenario", columns="year", values="spearman_vs_baseline").reindex([k for k in keys if k != "baseline"])
m.insert(0, "Specification", [labels[k] for k in m.index])
m = m.reset_index(drop=True)
m.columns = ["Specification"] + [str(c) for c in m.columns[1:]]
sheet_title(ws, "Sensitivity", "Spearman correlation between each specification and the baseline ranking, by year. 1.00 = identical ranking.", len(m.columns) + 2)
end = write_table(ws, m, 5, None, {i: "0.000" for i in range(2, len(m.columns) + 1)}, widths=[40] + [9] * (len(m.columns) - 1), zebra=False)
lc = get_column_letter(len(m.columns))
ws.conditional_formatting.add(f"B6:{lc}{end}", ColorScaleRule(start_type="num", start_value=0.75, start_color="FBE3B0", end_type="num", end_value=1, end_color="FFFFFF"))
c = ws.cell(5, len(m.columns) + 1, "Lowest"); c.fill = fill(NAVY); c.font = font(9, True, "FFFFFF")
for rr in range(6, end + 1):
    cc = ws.cell(rr, len(m.columns) + 1, f"=MIN(B{rr}:{lc}{rr})"); cc.number_format = "0.000"; cc.font = font(9, True, NAVY)
note(ws, end + 2, "Read: the digital-heavy row stays close to 1 because only three states have documented cloud regions, so extra weight on that pillar barely reorders the rest. The geometric-mean and vulnerability-heavy rows move rankings the most.", len(m.columns) + 1)

# ================================================================== Coverage
ws = wb.create_sheet("Coverage")
cv = cov.copy()
cv.insert(0, "State", cv.territory_id.map(NAME))
cv = cv.rename(columns={"territory_id": "Code", "year": "Year", "coverage": "Coverage (0–1)", "coverage_W": "Water", "coverage_E": "Energy",
                        "coverage_D": "Digital", "coverage_V": "Vulnerability", "tier": "Tier", "n_known": "Known", "n_derived": "Derived",
                        "n_estimated": "Estimated", "n_modelled": "Modelled", "n_unknown": "Unknown"})
# Tier sits in column H (the Guide's formulas count it there)
cv = cv[["State", "Code", "Year", "Coverage (0–1)", "Water", "Energy", "Digital", "Tier", "Vulnerability", "Known", "Derived", "Estimated", "Modelled", "Unknown"]]
sheet_title(ws, "Coverage", "How much we actually know for each state-year. Each indicator weighs by how it was obtained and how recent it is.", 14)
note(ws, 3, "Weights: recorded 1.0 · geospatial 0.9 · derived 0.85 · documentary 0.7 · estimated 0.6 · modelled 0.5 · unknown 0; minus 0.15 per year of age after the first; ×0.75 partial year; ×0.7 undated. Counts on the right refer to the 8 index indicators.", 14)
cols = [NAVY] * 4 + [PILLAR["W"][0], PILLAR["E"][0], PILLAR["D"][0], NAVY, PILLAR["V"][0]] + ["4F86B0", "7A8797", "B9852A", "7D6CC9", "8C8C8C"]
end = write_table(ws, cv, 5, cols, {4: "0.00", 5: "0.00", 6: "0.00", 7: "0.00", 9: "0.00"}, widths=[20, 6, 7, 11, 9, 9, 9, 12, 12, 8, 8, 9, 9, 9])
for t, hx in TIER.items():
    ws.conditional_formatting.add(f"H6:H{end}", CellIsRule(operator="equal", formula=[f'"{t}"'], fill=fill(hx)))
ws.conditional_formatting.add(f"D6:D{end}", ColorScaleRule(start_type="num", start_value=0, start_color="FBE3B0", end_type="num", end_value=1, end_color="BFE3F5"))

# ================================================================== Variables
ws = wb.create_sheet("Variables")
rows = []
for k, v in VARS["variables"].items():
    rows.append({"Pillar": VARS["pillars"][v["pillar"]]["en"], "Variable id": k, "Name": v["label"]["en"], "Unit": v["unit"],
                 "In the index?": "yes" if v["in_index"] else "no", "Direction": {1: "higher = more pressure", -1: "higher = less pressure", 0: "descriptive"}[v["direction"]],
                 "Transform": v["transform"], "Usual status": v["epistemic_status"], "Sources": ", ".join(v["source_ids"]),
                 "Definition": v["definition"]["en"], "Nombre (ES)": v["label"]["es"]})
vd = pd.DataFrame(rows)
sheet_title(ws, "Variables", "Data dictionary. Every column in Indicators and every variable_id in Observations is defined here.", 11)
end = write_table(ws, vd, 5, None, None, wrap_cols=(10,), widths=[22, 26, 30, 20, 9, 18, 9, 13, 14, 70, 28], zebra=False)
for rr in range(6, end + 1):
    key = [k for k, v in VARS["pillars"].items() if v["en"] == ws.cell(rr, 1).value][0]
    ws.cell(rr, 1).fill = fill(PILLAR[key][1]); ws.cell(rr, 1).font = font(9, True, PILLAR[key][0])
    ws.cell(rr, 8).fill = fill(LEDGER[ws.cell(rr, 8).value])

# ================================================================== Sources
ws = wb.create_sheet("Sources")
src = pd.read_csv(PUB / "sources.csv", dtype=str).fillna("")
src = src[["id", "institution", "title", "dataset_type", "license", "redistribution", "publication_date", "access_date", "temporal_coverage", "url_landing", "notes"]]
src.columns = ["ID", "Institution", "Dataset / document", "Type", "Licence / terms", "What we publish", "Published", "Accessed", "Period covered", "Link", "Notes"]
src["What we publish"] = src["What we publish"].map({"raw_ok": "raw allowed (fetched by script)", "derived_only": "aggregates only", "metadata_only": "citation only"})
sheet_title(ws, "Sources", "Every source is registered before its data enter the pipeline. Publicly accessible does not mean freely redistributable.", 11)
end = write_table(ws, src, 5, None, None, wrap_cols=(2, 3, 5, 9, 11), widths=[6, 28, 44, 13, 26, 18, 12, 11, 22, 40, 50], zebra=False)
for rr in range(6, end + 1):
    ws.cell(rr, 4).fill = fill(LEDGER.get(ws.cell(rr, 4).value, "FFFFFF"))
    link = ws.cell(rr, 10)
    if link.value:
        link.hyperlink = link.value; link.font = font(9, color="2F6FA8", underline="single")

# ================================================================== Cloud register
ws = wb.create_sheet("Cloud register")
reg = pd.read_csv(PUB / "digital_infrastructure_register.csv", dtype=str).fillna("")
reg.insert(1, "State", reg.cve_ent.str.zfill(2).map(NAME))
reg = reg.drop(columns=["cve_ent", "infrastructure_type"]).rename(columns={"record_id": "Record", "announced_date": "Announced", "operational_date": "Operational",
                                                                           "date_precision": "Date precision", "availability_zones": "Availability zones",
                                                                           "location_confidence": "Location source", "source_ids": "Sources", "notes": "Notes"})
sheet_title(ws, "Cloud register", "Hyperscale public-cloud regions documented in Mexico. Operators and facility locations are deliberately not recorded.", 9, band="D")
end = write_table(ws, reg, 5, [PILLAR["D"][0]] * len(reg.columns), None, wrap_cols=(9,), widths=[9, 18, 11, 12, 11, 11, 12, 12, 60], zebra=False)
ctx = pd.read_csv(PUB / "infrastructure_context.csv", dtype=str).fillna("")
note(ws, end + 2, f"Context only, never in the index: {NAME[ctx.cve_ent[0].zfill(2)]}, {ctx.value[0]} {ctx.unit[0]} of data-centre capacity reported operating as of {ctx.reference_date[0]} (source {ctx.source_id[0]}, figure without a named primary source).", 9)

# ================================================================== Governance
ws = wb.create_sheet("Governance")
gv = pd.read_csv(PUB / "governance_instruments.csv", dtype=str).fillna("")
gv = gv[["instrument_id", "name_en", "instrument_type", "level", "domains", "category", "date", "source_id", "description_en", "open_question_en", "name_es"]]
gv.columns = ["ID", "Instrument", "Type", "Level", "Domains", "Category", "Date", "Source", "What it does (neutral description)", "Open question", "Nombre (ES)"]
sheet_title(ws, "Governance", "Federal legal and institutional context. Descriptive only: nothing here is a legal conclusion and nothing enters the index.", 11, band="G")
end = write_table(ws, gv, 5, [PILLAR["G"][0]] * 11, None, wrap_cols=(2, 9, 10, 11), widths=[5, 34, 11, 9, 16, 22, 11, 7, 60, 44, 34], zebra=False)
for rr in range(6, end + 1):
    if ws.cell(rr, 10).value:
        ws.cell(rr, 10).fill = fill("FCF3DD")
note(ws, end + 2, "Categories: applicable_framework · regulatory_constraint · institutional_competence · policy_priority · resource_allocation · legal_question · data_insufficient · requires_legal_analysis. Dates of the 2025 water reform and the 2024 constitutional energy reform are pending verification against the DOF.", 11)

# ================================================================== Quality checks
ws = wb.create_sheet("Quality checks")
qc = pd.DataFrame(QA["checks"])[["check", "status", "count", "detail"]]
qc.columns = ["Check", "Result", "Count", "Detail"]
s = QA["summary"]
sheet_title(ws, "Quality checks", f"{s['checks']['pass']} checks passed · {s['checks']['warn']} warnings · {s['checks']['fail']} failed · validated {s['generated']}", 4)
note(ws, 3, f"Index cells (state × year × indicator): {s['index_cells']['complete_pct']}% complete, {s['index_cells']['partial_pct']}% partial (carried forward, partial-year, undated or projected), {s['index_cells']['missing_pct']}% missing.", 4)
end = write_table(ws, qc, 5, None, None, wrap_cols=(4,), widths=[42, 9, 8, 110], zebra=False)
ws.conditional_formatting.add(f"B6:B{end}", CellIsRule(operator="equal", formula=['"pass"'], fill=fill("DDF1E6")))
ws.conditional_formatting.add(f"B6:B{end}", CellIsRule(operator="equal", formula=['"info"'], fill=fill("E6EBF1")))
ws.conditional_formatting.add(f"B6:B{end}", CellIsRule(operator="equal", formula=['"fail"'], fill=fill("F6D3CF")))

# ================================================================== Observations (long)
ws = wb.create_sheet("Observations")
ob = obs.copy()
ob.insert(3, "variable_name", ob.variable_id.map(lambda k: VARS["variables"][k]["label"]["en"]))
ob = ob[["territory_name", "territory_id", "year", "variable_id", "variable_name", "value", "unit", "epistemic_status", "source_id", "vintage_year", "reference_date", "qa_flags", "obs_id"]]
ob.columns = ["State", "Code", "Year", "Variable id", "Variable", "Value", "Unit", "Status", "Sources", "Vintage", "Refers to", "Flags", "Observation id"]
sheet_title(ws, "Observations", "The full dataset in long form: one row per state × year × variable. Blank value = unknown, never zero.", 13)
end = write_table(ws, ob, 5, None, {6: "#,##0.###", 10: "0"}, widths=[18, 6, 6, 26, 34, 12, 20, 13, 12, 8, 18, 36, 14], zebra=False)
for st, hx in LEDGER.items():
    ws.conditional_formatting.add(f"H6:H{end}", CellIsRule(operator="equal", formula=[f'"{st}"'], fill=fill(hx)))

# row counts and quick facts on the Guide (formulas over exact table ranges)
count_col = {"Variables": "B"}
for name, rr in ROWCOUNT_ROWS.items():
    lo, hi = ENDS[name]
    col = count_col.get(name, "A")
    c = g.cell(rr, 3, f"=COUNTA('{name}'!{col}{lo}:{col}{hi})"); c.font = font(9, False, MUTED); c.number_format = "#,##0"
s0, s1 = ENDS["SRTI by year"]; c0, c1 = ENDS["Coverage"]; k0, k1 = ENDS["Cloud register"]; q0, q1 = ENDS["Sources"]
rng = f"'SRTI by year'!{srti_col}{s0}:{srti_col}{s1}"
facts = [
    (f"Highest SRTI in {y_last}", f"=MAX({rng})", "0.0"),
    (f"State with the highest SRTI in {y_last}", f"=INDEX('SRTI by year'!A{s0}:A{s1},MATCH(MAX({rng}),{rng},0))", "@"),
    (f"Median SRTI in {y_last}", f"=MEDIAN({rng})", "0.0"),
    ("Territory-years with moderate coverage", f'=COUNTIF(Coverage!H{c0}:H{c1},"moderate")', "0"),
    ("Territory-years with limited coverage", f'=COUNTIF(Coverage!H{c0}:H{c1},"limited")', "0"),
    ("Documented cloud regions", f"=COUNTA('Cloud register'!A{k0}:A{k1})", "0"),
    ("Sources registered", f"=COUNTA(Sources!A{q0}:A{q1})", "0"),
]
assert len(facts) == FACT_LABELS
for i, (lab, f_, fmt) in enumerate(facts):
    rr = facts_row + i
    g.cell(rr, 1, lab).font = font(9, False, INK)
    c = g.cell(rr, 2, f_); c.font = font(10, True, NAVY); c.number_format = fmt; c.alignment = Alignment(horizontal="left")
    g.cell(rr, 1).border = BORDER; g.cell(rr, 2).border = BORDER

for w in wb.worksheets:
    w.sheet_properties.tabColor = {"Cloud register": PILLAR["D"][0], "Governance": PILLAR["G"][0], "Guide": "E8B04A"}.get(w.title, NAVY)
    w.page_setup.orientation = "landscape"

out = ROOT / "dist" / f"DIRG_Observatory_v{VERSION}.xlsx"
out.parent.mkdir(exist_ok=True)
wb.save(out)
print(out)
